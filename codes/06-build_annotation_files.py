# -*- coding: utf-8 -*-
"""
قدم ۰۶: ساختِ فایل‌های Excel برای حاشیه‌نویسیِ دستی

طرحِ نمونه‌گیری (گزینه‌ی «الف»):
  لیلیا-کودک ۳۵۰ | مینو-کودک ۴۰۰ | لیلیا-درونداد ۳۰۰ | مینو-درونداد ۳۰۰  = ۱٬۳۵۰
  - نمونه از مجموعه‌ی نشان‌دار + بی‌نشان «با هم» گرفته می‌شود تا نسبتِ واقعیِ «را» حفظ شود
  - نمونه در همه‌ی جلسه‌ها پخش است (متناسب با حجمِ هر جلسه) تا اثرِ سن قابل بررسی بماند
  - فقط مواردی که فعلشان پیدا شده وارد می‌شوند (در هر دو گروهِ نشان‌دار و بی‌نشان یکسان)

کدگذاریِ خودکار (فقط پیشنهاد؛ شما تأیید یا اصلاح می‌کنید):
  - نوعِ گروهِ اسمی از %xmor
  - معرفگی برای ضمیر، اشاره، اسمِ خاص، ضمیرِ ملکی، «این/اون»، «یه»
  - جانداری برای ضمایرِ اول و دوم شخص و «کی/چی»
  - احتمالِ تکرارِ گفته‌ی قبلی (تقلید)

ورودی:
  data/processed/03-utterances.csv
  data/processed/03-tokens.csv
  data/processed/04-marked_objects.csv
  data/processed/05-unmarked_candidates_filtered.csv

خروجی:
  data/annotation/06-annotation_main.xlsx     فایلِ اصلی (همه‌ی ۱٬۳۵۰ مورد + برگه‌ی واژه‌ها + راهنما)
  data/annotation/06-annotation_rezai.xlsx    فایلِ ارزیابِ دوم (۲۰٪ از همان موارد، بدون کدهای شما)
  data/processed/06-sample.csv                فهرستِ موارد نمونه (برای ادغامِ بعدی)

⚠️ این اسکریپت فایل‌های Excelِ موجود را بازنویسی نمی‌کند، تا حاشیه‌نویسیِ شما از بین نرود.
"""
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT_DIR = Path(__file__).resolve().parents[1]
PROC_DIR = PROJECT_DIR / "data" / "processed"
ANN_DIR = PROJECT_DIR / "data" / "annotation"
ANN_DIR.mkdir(parents=True, exist_ok=True)

SEED = 42
TARGETS = {("Lilia", "CHILD"): 350, ("Minu", "CHILD"): 400,
           ("Lilia", "INPUT"): 300, ("Minu", "INPUT"): 300}
REL_FRAC = 0.20          # سهمِ ارزیابِ دوم

MAIN_XLSX = ANN_DIR / "06-annotation_main.xlsx"
REL_XLSX = ANN_DIR / "06-annotation_rezai.xlsx"
for f in (MAIN_XLSX, REL_XLSX):
    if f.exists():
        raise SystemExit(f"❌ {f.name} از قبل وجود دارد. برای جلوگیری از پاک شدنِ حاشیه‌نویسی، "
                         f"اگر واقعاً می‌خواهید از نو بسازید، اول خودتان آن را پاک یا جابه‌جا کنید.")

# ------------------------------------------------------------
# خواندنِ داده
# ------------------------------------------------------------
utts = pd.read_csv(PROC_DIR / "03-utterances.csv", keep_default_na=False)
toks = pd.read_csv(PROC_DIR / "03-tokens.csv", keep_default_na=False)
marked = pd.read_csv(PROC_DIR / "04-marked_objects.csv", keep_default_na=False)
unmarked = pd.read_csv(PROC_DIR / "05-unmarked_candidates_filtered.csv", keep_default_na=False)

tok_index = toks.set_index(["utt_id", "position"])

# گفته‌ی قبلی و بعدی در همان فایل (برای بافت)
utts["prev_text"] = utts.groupby("file")["text"].shift(1).fillna("")
utts["prev_speaker"] = utts.groupby("file")["speaker"].shift(1).fillna("")
utts["next_text"] = utts.groupby("file")["text"].shift(-1).fillna("")
utts["next_speaker"] = utts.groupby("file")["speaker"].shift(-1).fillna("")
ctx = utts.set_index("utt_id")

# ------------------------------------------------------------
# مجموعه‌ی نمونه‌گیری
# ------------------------------------------------------------
marked_ok = marked[(marked["host_lemma"] != "") & (marked["host_pos"] != "unk")
                   & (marked["verb_key"] != "")].copy()

# همان محدودیت‌های فعلیِ قدم ۰۵ روی نشان‌دارها هم اعمال شود تا «محدوده‌ی تغییرپذیری»
# برای هر دو گروه یکسان باشد (مثلاً «تو رو فراموش ... بودن» که فعلش اشتباه پیدا شده).
# ⚠️ این دو مقدار باید با قدم ۰۵ یکی باشند.
MIN_RA = 3
NON_OBJECT_VERBS = {
    "budæn", "shodæn", "amædæn", "ræftæn", "mandæn", "neshæstæn",
    "oftadæn", "xabidæn", "istadæn", "dævidæn", "gæshtæn", "residæn",
}
ra_count = marked[marked["group"] != "EXCLUDE"].groupby("verb_key").size()
n_before = len(marked_ok)
marked_ok = marked_ok[(marked_ok["verb_key"].map(ra_count).fillna(0) >= MIN_RA)
                      & ~marked_ok["verb_key"].str.split("_").str[-1].isin(NON_OBJECT_VERBS)]
print(f"نشان‌دارهای حذف‌شده به دلیلِ فعلِ ربطی/ناگذر یا کم‌تکرار: {n_before - len(marked_ok)}")

pool = pd.concat([marked_ok, unmarked], ignore_index=True)
pool = pool[pool["group"].isin(["CHILD", "INPUT"])].copy()

samples = []
for (child, group), n_target in TARGETS.items():
    sub = pool[(pool["child"] == child) & (pool["group"] == group)]
    if len(sub) <= n_target:
        samples.append(sub)
        continue
    frac = n_target / len(sub)
    # متناسب با حجمِ هر جلسه
    s = sub.groupby("file").sample(frac=frac, random_state=SEED)
    diff = n_target - len(s)
    if diff > 0:
        s = pd.concat([s, sub.drop(s.index).sample(n=diff, random_state=SEED)])
    elif diff < 0:
        s = s.drop(s.sample(n=-diff, random_state=SEED).index)
    samples.append(s)

sample = pd.concat(samples).sample(frac=1, random_state=SEED).reset_index(drop=True)
sample.insert(0, "id", [f"A{i:04d}" for i in range(1, len(sample) + 1)])

# ------------------------------------------------------------
# پیشنهادهای خودکار
# ------------------------------------------------------------
PERSONAL = {"mæn", "to", "ma", "shoma", "u", "ishun", "ishan"}
HUMAN_PRON = {"mæn", "to", "ma", "shoma"}
DEM_DET = {"in", "un", "hæmin", "hæmun"}
POSSESSIVE = {"1S", "2S", "3S", "1P", "2P", "3P"}


def to_int(x):
    """جایگاه‌ها گاهی به‌صورتِ «1.0» ذخیره شده‌اند؛ به عددِ صحیح تبدیل می‌کند."""
    try:
        return int(float(x))
    except (ValueError, TypeError):
        return None


def get_tok(utt_id, pos):
    pos = to_int(pos)
    if pos is None or pos < 0:
        return None
    try:
        return tok_index.loc[(utt_id, pos)]
    except KeyError:
        return None


def np_type(r):
    p = r["host_pos"]
    if p.startswith("pro:dem"):
        return "اشاره"
    if p.startswith("pro"):
        return "ضمیر شخصی" if r["host_lemma"] in PERSONAL else "ضمیر (دیگر)"
    if p.startswith("n:prop"):
        return "اسم خاص"
    if p.startswith("wh"):
        return "پرسش‌واژه"
    if p.startswith(("qn", "num")):
        return "کمیت‌نما/عدد"
    if p.startswith("pv"):
        return "جزء فعل مرکب"
    return "اسم عام"


def definiteness_auto(r):
    t = np_type(r)
    if t in ("اشاره", "ضمیر شخصی", "اسم خاص"):
        return "معرفه"
    if t == "پرسش‌واژه":
        return "پرسشی"
    sufs = set(str(r["host_suffixes"]).split())
    if sufs & POSSESSIVE:
        return "معرفه"
    hp = to_int(r["host_position"])
    prev = get_tok(r["utt_id"], hp - 1) if hp is not None else None
    if prev is not None:
        if str(prev["lemma"]) in DEM_DET and str(prev["pos"]).startswith("pro:dem"):
            return "معرفه"
        if str(prev["lemma"]) in {"ye", "yek"}:
            return "نکره (نوعش را تعیین کنید)"
    if "INDEF" in sufs:
        return "نکره (نوعش را تعیین کنید)"
    return ""   # اسمِ عامِ بی‌نشانه: قضاوتِ شما لازم است


def animacy_auto(r):
    if r["host_lemma"] in HUMAN_PRON or r["host_lemma"] == "ki":
        return "انسان"
    if r["host_lemma"] == "chi":
        return "بی‌جان"
    if np_type(r) in ("اسم عام", "اسم خاص", "کمیت‌نما/عدد", "جزء فعل مرکب"):
        return "← برگه‌ی واژه‌ها"
    return ""   # اشاره و ضمیرِ سوم‌شخص: بسته به مرجع


def norm(s):
    return re.sub(r"[^\wæ]+", " ", str(s)).strip().lower()


def imitation_auto(r):
    c = ctx.loc[r["utt_id"]]
    if c["prev_speaker"] and c["prev_speaker"] != r["speaker"]:
        a, b = norm(c["text"]), norm(c["prev_text"])
        if a and (a == b or (len(a) > 3 and a in b)):
            return "احتمالاً تکرار"
    return ""


sample["np_type"] = sample.apply(np_type, axis=1)
sample["def_auto"] = sample.apply(definiteness_auto, axis=1)
sample["anim_auto"] = sample.apply(animacy_auto, axis=1)
sample["imit_auto"] = sample.apply(imitation_auto, axis=1)
for col in ["text", "prev_text", "prev_speaker", "next_text", "next_speaker"]:
    sample[col] = sample["utt_id"].map(ctx[col])
sample["reliability"] = 0
rel_idx = sample.groupby(["child", "group"]).sample(frac=REL_FRAC, random_state=SEED).index
sample.loc[rel_idx, "reliability"] = 1

sample.to_csv(PROC_DIR / "06-sample.csv", index=False, encoding="utf-8-sig")

# ------------------------------------------------------------
# ساختِ Excel
# ------------------------------------------------------------
YES_NO = '"بله,خیر,؟"'
DEF_LIST = '"معرفه,نکره‌ی مشخص,نکره‌ی نامشخص,جنس/عام,پرسشی,نامعلوم"'
ANIM_LIST = '"انسان,حیوان,بی‌جان,نامعلوم"'
IMIT_LIST = '"خودجوش,تکرار,قالب ثابت"'

HEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
AUTO_FILL = PatternFill("solid", fgColor="EDEDED")
MANUAL_FILL = PatternFill("solid", fgColor="FFF2CC")

# (عنوان، کلید در sample یا None برای ستونِ خالی، پهنا، نوع: info/auto/manual، فهرستِ کشویی)
COLUMNS = [
    ("شناسه", "id", 8, "info", None),
    ("کودک", "child", 7, "info", None),
    ("گوینده", "speaker", 7, "info", None),
    ("سن (ماه)", "age_months", 7, "info", None),
    ("گفته‌ی قبلی", "prev_text", 32, "info", None),
    ("گفته", "text", 40, "info", None),
    ("گفته‌ی بعدی", "next_text", 32, "info", None),
    ("اسم", "host_lemma", 12, "info", None),
    ("فعل", "verb_key", 13, "info", None),
    ("«را» دارد؟", "marked", 7, "auto", None),
    ("۱) مفعولِ مستقیمِ همین فعل است؟", None, 12, "manual", YES_NO),
    ("نوعِ گروهِ اسمی (خودکار)", "np_type", 12, "auto", None),
    ("معرفگی (پیشنهادِ خودکار)", "def_auto", 14, "auto", None),
    ("۲) معرفگی (نهایی)", None, 14, "manual", DEF_LIST),
    ("جانداری (پیشنهادِ خودکار)", "anim_auto", 13, "auto", None),
    ("۳) جانداری (فقط اشاره/ضمیر)", None, 13, "manual", ANIM_LIST),
    ("تقلید؟ (پیشنهادِ خودکار)", "imit_auto", 11, "auto", None),
    ("۴) خودجوش/تکرار/قالب", None, 11, "manual", IMIT_LIST),
    ("توضیح", None, 25, "manual", None),
]


def write_items(ws, df, blind_marked_default=False):
    ws.sheet_view.rightToLeft = True
    for j, (title, _, width, kind, _) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=j, value=title)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 45
    for i, (_, r) in enumerate(df.iterrows(), 2):
        for j, (title, key, _, kind, _) in enumerate(COLUMNS, 1):
            val = r[key] if key else ""
            if key == "marked":
                val = "بله" if int(r["marked"]) == 1 else "خیر"
            if title.startswith("۱)") and int(r["marked"]) == 1:
                val = "بله"          # مفعولِ نشان‌دار قطعاً مفعول است
            c = ws.cell(row=i, column=j, value=val)
            c.alignment = Alignment(wrap_text=True, vertical="top",
                                    horizontal="left" if key in ("prev_text", "text", "next_text") else "center")
            if kind == "auto":
                c.fill = AUTO_FILL
            elif kind == "manual":
                c.fill = MANUAL_FILL
    n = len(df) + 1
    for j, (_, _, _, kind, options) in enumerate(COLUMNS, 1):
        if options:
            letter = ws.cell(row=1, column=j).column_letter
            dv = DataValidation(type="list", formula1=options, allow_blank=True)
            dv.add(f"{letter}2:{letter}{n}")
            ws.add_data_validation(dv)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUMNS)).column_letter}{n}"


def write_lemmas(ws, df):
    """برگه‌ی واژه‌ها: جانداری یک بار برای هر واژه."""
    ws.sheet_view.rightToLeft = True
    lem = df[df["anim_auto"] == "← برگه‌ی واژه‌ها"]
    table = (lem.groupby("host_lemma")
             .agg(count=("id", "size"), example=("text", "first"))
             .sort_values("count", ascending=False).reset_index())
    heads = [("واژه", 16), ("تعداد", 7), ("نمونه", 50), ("جانداری", 12), ("توضیح", 25)]
    for j, (h, w) in enumerate(heads, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = Font(bold=True), HEAD_FILL
        ws.column_dimensions[c.column_letter].width = w
    for i, r in enumerate(table.itertuples(), 2):
        ws.cell(row=i, column=1, value=r.host_lemma)
        ws.cell(row=i, column=2, value=r.count)
        ws.cell(row=i, column=3, value=r.example).alignment = Alignment(horizontal="left")
        ws.cell(row=i, column=4).fill = MANUAL_FILL
        ws.cell(row=i, column=5).fill = MANUAL_FILL
    dv = DataValidation(type="list", formula1=ANIM_LIST, allow_blank=True)
    dv.add(f"D2:D{len(table) + 1}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "A2"
    return len(table)


GUIDE = [
    "راهنمای حاشیه‌نویسی",
    "",
    "فقط ستون‌های زرد را پر کنید. ستون‌های خاکستری پیشنهادِ خودکارند؛ اگر درست‌اند، همان را در ستونِ زردِ کنارش انتخاب کنید.",
    "برای هر ستونِ زرد یک فهرستِ کشویی هست.",
    "",
    "۱) مفعولِ مستقیمِ همین فعل است؟",
    "   آزمون: پس از اسم «رو» بگذارید. اگر اسم همچنان مفعولِ همان فعل بماند → بله. اگر نقشش عوض شود یا جمله نادرست شود → خیر.",
    "   خیر: فاعل («بچه‌ها می‌گن»)، مخاطب («لیلیا بخون»)، تأکید بر فاعل («خودت بخوری»)، قید («یه ذره»، «دیگه»)، درونِ گروهِ حرف‌اضافه‌ای.",
    "   برای موارد «را»دار از پیش «بله» گذاشته شده. اگر خیر شد، ستون‌های بعدی را خالی بگذارید.",
    "",
    "۲) معرفگی (نهایی) — نقشِ این اسم در همین گفته:",
    "   معرفه: شنونده دقیقاً می‌داند کدام است (کیکِ روی میز، دستم، این، اسم خاص).",
    "   نکره‌ی مشخص: یک چیزِ معین که شنونده نمی‌شناسد («یه کتاب خریدم که جلدش قرمزه»).",
    "   نکره‌ی نامشخص: هر نمونه‌ای از آن نوع («یه لیوان آب بیار»، «کیک می‌خوری؟»).",
    "   جنس/عام: کلِ نوع، نه یک نمونه («کارتون دیدن بده»، «من جوجه دوست دارم»).",
    "   پرسشی: «چی»، «کدوم».  نامعلوم: اگر بافت کافی نیست.",
    "",
    "۳) جانداری — فقط برای اشاره و ضمیرِ سوم‌شخص (این، اون، همه)، بسته به اینکه در این گفته به چه اشاره دارد.",
    "   برای اسم‌ها جانداری را یک بار در برگه‌ی «واژه‌ها» بنویسید.",
    "",
    "۴) خودجوش / تکرار / قالب:",
    "   تکرار: گوینده گفته‌ی قبلیِ دیگری را تکرار کرده.  قالب ثابت: عبارتِ حفظ‌شده (مثلاً «اینو بده» به‌صورتِ یک تکه).",
    "   خودجوش: هیچ‌کدام. (این ستون مهم‌تر از همه برای گفتارِ کودک است.)",
    "",
    "نکته: ستون‌های «گفته‌ی قبلی» و «گفته‌ی بعدی» برای فهمِ بافت‌اند.",
]


def write_guide(ws):
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 130
    for i, line in enumerate(GUIDE, 1):
        c = ws.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True)
        if i == 1 or line[:2] in ("۱)", "۲)", "۳)", "۴)"):
            c.font = Font(bold=True, size=12 if i > 1 else 14)


def build(path, df):
    wb = Workbook()
    write_guide(wb.active)
    wb.active.title = "راهنما"
    write_items(wb.create_sheet("موارد"), df)
    n_lem = write_lemmas(wb.create_sheet("واژه‌ها"), df)
    wb.save(path)
    return n_lem


n_lem_main = build(MAIN_XLSX, sample)
rel = sample[sample["reliability"] == 1].sample(frac=1, random_state=SEED + 1)
n_lem_rel = build(REL_XLSX, rel)

# ------------------------------------------------------------
# گزارش
# ------------------------------------------------------------
print("تعدادِ نمونه به تفکیکِ گروه (نشان‌دار / بی‌نشان):")
print(sample.groupby(["child", "group"])["marked"].agg(total="size", marked="sum").to_string())
print(f"\nنیازمندِ تأییدِ «مفعول است؟» (بی‌نشان): {(sample['marked'] == 0).sum()}")
print(f"نیازمندِ قضاوتِ معرفگی (بدون پیشنهاد): {(sample['def_auto'] == '').sum()}")
print(f"واژه‌های برگه‌ی «واژه‌ها» (جانداری): {n_lem_main}")
print(f"جانداریِ تک‌موردی (اشاره/ضمیر): {(sample['anim_auto'] == '').sum()}")
print(f"\nفایلِ ارزیابِ دوم: {len(rel)} مورد، {n_lem_rel} واژه")
print(f"\n✅ {MAIN_XLSX.name} و {REL_XLSX.name} در {ANN_DIR} ساخته شد.")
